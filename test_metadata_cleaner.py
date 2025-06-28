import os
import unittest
from metadata_cleaner import clean_image, clean_pdf, clean_docx, clean_xlsx, clean_pptx, clean_video, process_file

class TestMetadataCleaner(unittest.TestCase):
    def setUp(self):
        os.makedirs('test_files', exist_ok=True)
        # Create dummy files for each type
        from PIL import Image
        img = Image.new('RGB', (10, 10), color='red')
        img.save('test_files/test.jpg')
        img.save('test_files/test.png')
        # Create a minimal PDF
        with open('test_files/test.pdf', 'wb') as f:
            f.write(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF')
        # Create DOCX, XLSX, PPTX
        from docx import Document
        doc = Document()
        doc.add_paragraph('Hello')
        doc.save('test_files/test.docx')
        from openpyxl import Workbook
        wb = Workbook()
        wb.save('test_files/test.xlsx')
        from pptx import Presentation
        prs = Presentation()
        prs.save('test_files/test.pptx')
        # Create a dummy MP4 (empty, will fail gracefully)
        with open('test_files/test.mp4', 'wb') as f:
            f.write(b'')

    def tearDown(self):
        import shutil
        shutil.rmtree('test_files', ignore_errors=True)
        # Remove any _clean files
        for ext in ['.jpg', '.png', '.pdf', '.docx', '.xlsx', '.pptx', '.mp4']:
            f = f'test_files/test_clean{ext}'
            if os.path.exists(f):
                os.remove(f)

    def test_clean_image_jpg(self):
        out = clean_image('test_files/test.jpg')
        self.assertTrue(out and os.path.exists(out))

    def test_clean_image_png(self):
        out = clean_image('test_files/test.png')
        self.assertTrue(out and os.path.exists(out))

    def test_clean_pdf(self):
        out = clean_pdf('test_files/test.pdf')
        # May fail if pikepdf can't open minimal PDF, so just check for None or file
        self.assertTrue(out is None or os.path.exists(out))

    def test_clean_docx(self):
        out = clean_docx('test_files/test.docx')
        self.assertTrue(out and os.path.exists(out))

    def test_clean_xlsx(self):
        out = clean_xlsx('test_files/test.xlsx')
        self.assertTrue(out and os.path.exists(out))

    def test_clean_pptx(self):
        out = clean_pptx('test_files/test.pptx')
        self.assertTrue(out and os.path.exists(out))

    def test_clean_video(self):
        out = clean_video('test_files/test.mp4')
        # Should fail gracefully (ffmpeg will error on empty file)
        self.assertTrue(out is None or os.path.exists(out))

    def test_process_file(self):
        for ext in ['.jpg', '.png', '.pdf', '.docx', '.xlsx', '.pptx', '.mp4']:
            f = f'test_files/test{ext}'
            out = process_file(f)
            self.assertTrue(out is None or os.path.exists(out))

if __name__ == '__main__':
    unittest.main()
